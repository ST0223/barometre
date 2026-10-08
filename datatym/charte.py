"""Charte documentaire Datatym appliquée aux sorties Word (logos Datatym et Agiltym en tête).

Tous les documents sont décrits par un « modèle de document » (dict sérialisable) :
    {titre, surtitre, sous_titre, meta[], en_tete, paysage, sommaire,
     sections: [{id, titre, niveau, source, blocs: [...]}]}
Blocs : p, liste, tableau, image, chiffres, saut.
Le même modèle sert à l'aperçu HTML et à l'édition dans l'interface.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from .config import BLEU, BLEU_CLAIR, GRIS, LOGO, LOGO_AGILTYM, MARINE, POLICE, BASE_DIR

GRAS_RE = re.compile(r"\*\*(.+?)\*\*")


def _rgb(hexa):
    return RGBColor.from_string(hexa)


def _police_run(run, taille=None, gras=None, couleur=None, italique=None):
    run.font.name = POLICE
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), POLICE)
    if taille:
        run.font.size = Pt(taille)
    if gras is not None:
        run.font.bold = gras
    if italique is not None:
        run.font.italic = italique
    if couleur:
        run.font.color.rgb = _rgb(couleur)


def _ombrer(cellule, hexa):
    tcpr = cellule._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tcpr.append(shd)


def _bordures_tableau(table, couleur="C9D3E0"):
    tblpr = table._element.tblPr
    bordures = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), couleur)
        bordures.append(el)
    tblpr.append(bordures)


def _repeter_entete(ligne):
    trpr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _champ(paragraphe, code, taille=9, couleur="6B7280"):
    run = paragraphe.add_run()
    _police_run(run, taille, couleur=couleur)
    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for el in (debut, instr, sep, txt, fin):
        run._element.append(el)


def ajouter_texte(paragraphe, texte, taille=None, couleur=None, gras_defaut=False, italique=None):
    """Ajoute du texte en interprétant **gras**."""
    morceaux = GRAS_RE.split(texte or "")
    for i, morceau in enumerate(morceaux):
        if not morceau:
            continue
        run = paragraphe.add_run(morceau)
        _police_run(run, taille, gras=gras_defaut or i % 2 == 1, couleur=couleur, italique=italique)


class DocDatatym:
    """Document Word préconfiguré : Times New Roman 12, interligne 1,5, justifié, logo Datatym."""

    def __init__(self, en_tete="", paysage=False, pied="Document interne — Confidentiel"):
        self.pied = pied
        self.doc = Document()
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        if paysage:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21)
        for marge in ("left_margin", "right_margin"):
            setattr(sec, marge, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2)
        self.largeur_utile = (sec.page_width - sec.left_margin - sec.right_margin)
        self._entete_pied(sec, en_tete)
        self.n_tableaux = 0
        self.n_figures = 0

    # ---------- styles ----------
    def _styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = POLICE
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(6)
        for nom, taille, couleur, avant in (("Heading 1", 16, MARINE, 18), ("Heading 2", 13.5, BLEU, 12), ("Heading 3", 12, MARINE, 8)):
            st = styles[nom]
            st.font.name = POLICE
            st.font.size = Pt(taille)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = _rgb(couleur)
            rpr = st.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(att), POLICE)
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(att)) is not None:
                    del rfonts.attrib[qn(att)]
            st.paragraph_format.space_before = Pt(avant)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.keep_with_next = True
        lp = styles["List Paragraph"] if "List Paragraph" in [s.name for s in styles] else None
        if lp is not None:
            lp.font.name = POLICE

    def logos(self, paragraphe, h_datatym, h_agiltym):
        """Datatym à gauche, Agiltym à droite, sur une même ligne."""
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.LEFT
        taquets = paragraphe.paragraph_format.tab_stops
        for pos in (4680, 9360):  # taquets hérités du style « En-tête » de Word
            taquets.add_tab_stop(Twips(pos), WD_TAB_ALIGNMENT.CLEAR)
        taquets.add_tab_stop(self.largeur_utile, WD_TAB_ALIGNMENT.RIGHT)
        if LOGO.exists():
            paragraphe.add_run().add_picture(str(LOGO), height=Cm(h_datatym))
        if LOGO_AGILTYM.exists():
            paragraphe.add_run("	").add_picture(str(LOGO_AGILTYM), height=Cm(h_agiltym))

    def _entete_pied(self, sec, en_tete):
        header = sec.header
        p = header.paragraphs[0]
        self.logos(p, 0.62, 0.52)
        p.paragraph_format.space_after = Pt(2)
        en_tete = en_tete or ""
        if len(en_tete) > 90:
            debut, _, fin = en_tete.rpartition(" — ")
            if not debut:
                debut, fin = en_tete, ""
            reste = max(20, 86 - len(fin))
            debut = debut[:reste].rsplit(" ", 1)[0] + "…" if len(debut) > reste else debut
            en_tete = f"{debut} — {fin}" if fin else debut
        p2 = header.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(0)
        _police_run(p2.add_run(en_tete), 8, couleur="6B7280")
        bordure = OxmlElement("w:pBdr")
        bas = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", BLEU)):
            bas.set(qn(k), v)
        bordure.append(bas)
        p2._p.get_or_add_pPr().append(bordure)

        for pied in (sec.footer, sec.first_page_footer):
            fp = pied.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = fp.add_run(f"{self.pied}   ·   Page ")
            _police_run(r, 9, couleur="6B7280")
            _champ(fp, "PAGE")
            r = fp.add_run(" / ")
            _police_run(r, 9, couleur="6B7280")
            _champ(fp, "NUMPAGES")

    # ---------- blocs ----------
    def couverture(self, surtitre, titre, sous_titre="", meta=()):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()
        if surtitre:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ajouter_texte(p, surtitre.upper(), 10.5, BLEU, gras_defaut=True)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.1
        ajouter_texte(p, titre, 26, MARINE, gras_defaut=True)
        if sous_titre:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ajouter_texte(p, sous_titre, 15, MARINE)
        p = d.add_paragraph()
        bordure = OxmlElement("w:pBdr")
        bas = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "1"), ("w:color", BLEU)):
            bas.set(qn(k), v)
        bordure.append(bas)
        p._p.get_or_add_pPr().append(bordure)
        for ligne in meta or ():
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            ajouter_texte(p, ligne, 11, "4B5563")
        self.saut_page()

    def sommaire(self):
        self.titre("Sommaire", 1)
        p = self.doc.add_paragraph()
        _champ(p, 'TOC \\o "1-2" \\h \\z \\u', 11, "000000")
        note = self.doc.add_paragraph()
        ajouter_texte(note, "Sous Word : clic droit sur le sommaire › Mettre à jour les champs.", 9, "6B7280", italique=True)
        self.saut_page()

    def titre(self, texte, niveau=1):
        self.doc.add_heading(texte, level=min(max(int(niveau or 1), 1), 3))

    def paragraphe(self, texte, style="normal"):
        p = self.doc.add_paragraph()
        if style == "note":
            ajouter_texte(p, texte, 10, "4B5563", italique=True)
        elif style == "encadre":
            ajouter_texte(p, texte, 12, MARINE)
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), BLEU_CLAIR)
            ppr.append(shd)
            bordure = OxmlElement("w:pBdr")
            gauche = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "24"), ("w:space", "8"), ("w:color", BLEU)):
                gauche.set(qn(k), v)
            bordure.append(gauche)
            ppr.append(bordure)
        else:
            ajouter_texte(p, texte, gras_defaut=(style == "gras"))
        return p

    def liste(self, items, numerotee=False):
        style = "List Number" if numerotee else "List Bullet"
        for item in items:
            p = self.doc.add_paragraph(style=style)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            ajouter_texte(p, item)

    def tableau(self, colonnes, lignes, legende="", largeurs=None, taille=10):
        if legende:
            self.n_tableaux += 1
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            ajouter_texte(p, f"Tableau {self.n_tableaux}. {legende}", 10, MARINE, gras_defaut=True)
        n = max(len(colonnes), max((len(l) for l in lignes), default=0), 1)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures_tableau(t)
        entete = t.rows[0]
        _repeter_entete(entete)
        for i in range(n):
            c = entete.cells[i]
            c.text = ""
            _ombrer(c, MARINE)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            ajouter_texte(p, str(colonnes[i]) if i < len(colonnes) else "", taille, "FFFFFF", gras_defaut=True)
        for k, ligne in enumerate(lignes):
            cells = t.add_row().cells
            for i in range(n):
                val = ligne[i] if i < len(ligne) else ""
                c = cells[i]
                if k % 2 == 1:
                    _ombrer(c, GRIS)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                for j, morceau in enumerate(str(val if val is not None else "").split("\n")):
                    if j:
                        p = c.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.line_spacing = 1.0
                        p.paragraph_format.space_after = Pt(1)
                    ajouter_texte(p, morceau, taille)
        if largeurs:
            total = sum(largeurs)
            for ligne in t.rows:
                for i, w in enumerate(largeurs[:n]):
                    ligne.cells[i].width = int(self.largeur_utile * w / total)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def chiffres(self, items):
        """Bandeau de chiffres clés (valeur + libellé)."""
        items = items[:4] if len(items) <= 4 else items
        n = min(len(items), 4) or 1
        for debut in range(0, len(items), n):
            t = self.doc.add_table(rows=2, cols=n)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, it in enumerate(items[debut:debut + n]):
                for r in (0, 1):
                    _ombrer(t.rows[r].cells[i], BLEU_CLAIR)
                p = t.rows[0].cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ajouter_texte(p, str(it.get("valeur", "")), 22, BLEU, gras_defaut=True)
                p = t.rows[1].cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                ajouter_texte(p, str(it.get("libelle", "")), 9.5, MARINE)
            self.doc.add_paragraph()

    def image(self, chemin, largeur_cm=15, legende=""):
        chemin = Path(chemin)
        if not chemin.is_absolute():
            chemin = BASE_DIR / chemin
        if not chemin.exists():
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(chemin), width=Cm(largeur_cm))
        if legende:
            self.n_figures += 1
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ajouter_texte(p, f"Graphique {self.n_figures}. {legende}", 9.5, "4B5563", italique=True)

    def saut_page(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def proprietes(self, titre="", sujet="", auteur="DATATYM"):
        cp = self.doc.core_properties
        cp.author = cp.last_modified_by = auteur
        cp.title, cp.subject = titre, sujet
        cp.keywords = "Datatym; baromètre; confidentiel"

    def enregistrer(self, chemin):
        Path(chemin).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(chemin))
        return chemin


def rendre_modele(modele, chemin):
    """Produit le .docx d'un modèle de document."""
    d = DocDatatym(modele.get("en_tete", ""), paysage=modele.get("paysage", False),
                   pied=modele.get("pied", "Document interne — Confidentiel"))
    d.proprietes(modele.get("titre", ""), modele.get("sous_titre", ""), modele.get("auteur", "DATATYM"))
    d.couverture(modele.get("surtitre", ""), modele.get("titre", ""), modele.get("sous_titre", ""), modele.get("meta", []))
    if modele.get("sommaire"):
        d.sommaire()
    for sec in modele.get("sections", []):
        if sec.get("titre"):
            d.titre(sec["titre"], sec.get("niveau", 1))
        for b in sec.get("blocs", []):
            t = b.get("type")
            if t == "p":
                d.paragraphe(b.get("texte", ""), b.get("style", "normal"))
            elif t == "liste":
                d.liste([i for i in b.get("items", []) if str(i).strip()], b.get("numerotee", False))
            elif t == "tableau":
                d.tableau(b.get("colonnes", []), b.get("lignes", []), b.get("legende", ""), b.get("largeurs"), b.get("taille", 10))
            elif t == "chiffres":
                d.chiffres(b.get("items", []))
            elif t == "image":
                d.image(b.get("chemin", ""), b.get("largeur_cm", 15), b.get("legende", ""))
            elif t == "saut":
                d.saut_page()
    return d.enregistrer(chemin)


# ---------- aides de construction de modèles ----------
def section(id_, titre, blocs, niveau=1, source="texte"):
    return {"id": id_, "titre": titre, "niveau": niveau, "source": source, "blocs": blocs}


def p(texte, style="normal"):
    return {"type": "p", "texte": texte, "style": style}


def liste(items, numerotee=False):
    return {"type": "liste", "items": list(items), "numerotee": numerotee}


def tableau(colonnes, lignes, legende="", largeurs=None, taille=10):
    b = {"type": "tableau", "colonnes": list(colonnes), "lignes": [list(map(lambda x: "" if x is None else str(x), l)) for l in lignes], "legende": legende, "taille": taille}
    if largeurs:
        b["largeurs"] = largeurs
    return b


def chiffres(items):
    return {"type": "chiffres", "items": items}


def image(chemin, legende="", largeur_cm=15):
    return {"type": "image", "chemin": str(chemin), "legende": legende, "largeur_cm": largeur_cm}
