"""Rappels de contrôle (back-checks) : un échantillon de répondants est rappelé pour vérifier que l'entretien a bien eu lieu
et que quelques réponses concordent. Tirage par enquêteur, fiche de rappel, résultats, enquêteurs à surveiller.

Un entretien déclaré « non réalisé » au rappel est retiré de la base : il ne compte plus ni dans les résultats ni dans la paie.
"""
import hashlib
import math

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .config import auteur_fichiers
from .formulaires import _feuille
from .outils import date_courte
from .store import maintenant as maintenant_iso

STATUTS = {
    "a_faire": "À faire",
    "confirme": "Confirmé",
    "ecart": "Confirmé avec écarts",
    "injoignable": "Injoignable",
    "refus": "Refuse le rappel",
    "non_realise": "Entretien non réalisé",
}
JOINTS = ("confirme", "ecart", "non_realise")      # le répondant a pu être interrogé
CONTACT = ("cl_rappel", "cl_entreprise", "cl_telephone")


def liste(projet):
    return (projet["etapes"]["collecte"].get("data") or {}).setdefault("rappels", [])


def invalides(projet):
    """Identifiants des entretiens déclarés non réalisés au rappel (exclus de la base)."""
    return {r["id"] for r in (projet["etapes"]["collecte"].get("data") or {}).get("rappels") or [] if r["statut"] == "non_realise"}


def _quest(projet):
    return projet["etapes"]["questionnaire"].get("data") or {}


def contact_disponible(projet):
    ids = {q["id"] for q in _quest(projet).get("questions_plates") or []}
    return "cl_telephone" in ids


def _part(projet):
    b = (projet["etapes"]["echantillonnage"].get("data") or {}).get("budget") or {}
    return float((b.get("hypotheses") or {}).get("part_rappels", 10) or 10)


def _ordre(projet, ident):
    """Ordre de tirage pseudo-aléatoire mais stable : le même entretien garde sa place d'un tirage à l'autre."""
    return hashlib.sha1(f"{projet['slug']}:{ident}".encode()).hexdigest()


def tirer(projet):
    """Ajoute les entretiens à rappeler pour atteindre la part prévue, enquêteur par enquêteur (au moins un chacun).
    Les entretiens dont le répondant a accepté le rappel et laissé un numéro passent en premier."""
    from .etapes import collecte
    quest = _quest(projet)
    df = collecte.lire(projet)
    if df is None or df.empty:
        raise ValueError("Aucune réponse importée : rien à rappeler pour l'instant.")
    controle = projet["etapes"]["collecte"]["data"].get("controle_qualite", True)
    cons = df.get("consentement", df.iloc[:, 0].map(lambda _: "oui")).astype(str).str.lower().isin(["oui", "1", "yes"])
    v = collecte.valides(df, quest) if controle else df[cons]
    v_enq, v_zone, v_prof = collecte.variables_cles(quest)
    col = "enqueteur" if "enqueteur" in v else v_enq
    par = {q["id"]: q for q in quest["questions_plates"]}
    var = {k: par[k]["var"] for k in par}
    lib = lambda qid, code: next((m["label"] for m in (par.get(qid) or {}).get("modalites", []) if m["code"] == str(code)), str(code or ""))  # noqa: E731
    deja = {r["id"] for r in liste(projet)}
    part = _part(projet)
    # questions à revérifier au téléphone : fonction, secteur et deux questions factuelles d'identification
    verif_ids = [k for k in ("id_profil", "id_secteur") if k in par]
    verif_ids += [q["id"] for q in quest["questions_plates"] if q.get("indicateur") and q["section"] == "A" and q["type"] == "select_one"][:2]

    nouveaux = []
    groupes = v.groupby(v[col].fillna("").astype(str).str.strip()) if col in v else [("", v)]
    for collecteur, g in groupes:
        cible = max(1, math.ceil(len(g) * part / 100))
        faits = sum(1 for r in liste(projet) if r["collecteur"] == (collecteur or "Non renseigné"))
        if faits >= cible:
            continue
        candidats = []
        for _, l in g.iterrows():
            ident = str(l.get("_id", ""))
            if not ident or ident in deja:
                continue
            accord = str(l.get(var.get("cl_rappel", "_"), "")) == "oui"
            tel = str(l.get(var.get("cl_telephone", "_"), "") or "").strip()
            candidats.append((0 if accord and tel else 1, _ordre(projet, ident), l))
        for _, _, l in sorted(candidats, key=lambda x: x[:2])[:cible - faits]:
            ident = str(l["_id"])
            nouveaux.append({
                "id": ident, "tire_le": maintenant_iso(), "collecteur": collecteur or "Non renseigné",
                "zone": lib("id_zone", l.get(v_zone, "")), "profil": lib("id_profil", l.get(v_prof, "")),
                "date": str(l.get("start", ""))[:10], "duree": str(l.get("duree_minutes", "") or ""),
                "contact": {"accord": str(l.get(var.get("cl_rappel", "_"), "") or ""),
                            "entreprise": str(l.get(var.get("cl_entreprise", "_"), "") or ""),
                            "telephone": str(l.get(var.get("cl_telephone", "_"), "") or "")},
                "verifs": [{"question": par[k]["texte"], "reponse": lib(k, l.get(var[k], ""))} for k in verif_ids
                           if var.get(k) in l and str(l.get(var[k], "")).strip()],
                "statut": "a_faire", "commentaire": "", "fait_le": None,
            })
    liste(projet).extend(nouveaux)
    return nouveaux


def maj(projet, ident, statut, commentaire=None):
    if statut not in STATUTS:
        raise ValueError("Résultat de rappel inconnu.")
    r = next((x for x in liste(projet) if x["id"] == str(ident)), None)
    if not r:
        raise ValueError("Rappel introuvable.")
    avant = r["statut"]
    r["statut"] = statut
    if commentaire is not None:
        r["commentaire"] = str(commentaire)[:500]
    r["fait_le"] = None if statut == "a_faire" else (r.get("fait_le") if avant == statut else maintenant_iso())
    return "non_realise" in (avant, statut) and avant != statut  # la base de réponses change


def synthese(projet):
    rs = liste(projet)
    par = {}
    for r in rs:
        s = par.setdefault(r["collecteur"], {"collecteur": r["collecteur"], "tires": 0, "faits": 0, **{k: 0 for k in STATUTS}})
        s["tires"] += 1
        s[r["statut"]] += 1
        s["faits"] += r["statut"] != "a_faire"
    lignes = []
    for s in par.values():
        joints = sum(s[k] for k in JOINTS)
        s["taux_confirmation"] = round(100 * (s["confirme"] + s["ecart"]) / joints, 1) if joints else None
        # un entretien inventé suffit ; des écarts répétés aussi, dès que trois répondants ont été joints
        s["alerte"] = s["non_realise"] > 0 or (joints >= 3 and s["ecart"] / joints >= 0.3)
        lignes.append(s)
    lignes.sort(key=lambda s: (not s["alerte"], s["collecteur"]))
    tot = {k: sum(s[k] for s in lignes) for k in ("tires", "faits", *STATUTS)}
    joints = sum(tot[k] for k in JOINTS)
    tot["taux_confirmation"] = round(100 * (tot["confirme"] + tot["ecart"]) / joints, 1) if joints else None
    return {"lignes": lignes, "total": tot, "part": _part(projet), "contact": contact_disponible(projet),
            "alertes": [s["collecteur"] for s in lignes if s["alerte"]]}


def xlsx(chemin, projet):
    """Fiche de rappel : un rappel par ligne, avec les réponses à revérifier et les colonnes à remplir."""
    rs = liste(projet)
    nb_v = max((len(r["verifs"]) for r in rs), default=0)
    cols = ["ID réponse", "Date", "Enquêteur", "Pays", "Fonction", "Entreprise", "Téléphone", "Accord rappel"]
    for k in range(nb_v):
        cols += [f"Question {k + 1}", f"Réponse enregistrée {k + 1}"]
    cols += ["Résultat", "Commentaire"]
    lignes = []
    for r in rs:
        l = [r["id"], date_courte(r["date"]) if r["date"] else "", r["collecteur"], r["zone"], r["profil"],
             r["contact"].get("entreprise", ""), r["contact"].get("telephone", ""), "Oui" if r["contact"].get("accord") == "oui" else ""]
        for k in range(nb_v):
            x = r["verifs"][k] if k < len(r["verifs"]) else {"question": "", "reponse": ""}
            l += [x["question"], x["reponse"]]
        lignes.append(l + [STATUTS[r["statut"]], r.get("commentaire", "")])
    wb = Workbook()
    ws = wb.active
    ws.title = "Rappels"
    _feuille(ws, cols, lignes, {"Entreprise": 22, "Téléphone": 18, "Résultat": 22, "Commentaire": 40})
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS.values()) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    col_res = len(cols) - 1
    from openpyxl.utils import get_column_letter
    dv.add(f"{get_column_letter(col_res)}2:{get_column_letter(col_res)}{max(len(lignes) + 1, 2)}")
    ws2 = wb.create_sheet("Script d'appel")
    _feuille(ws2, ["Étape", "Ce qu'on dit ou fait"], [
        ["1. Présentation", "Bonjour, je vous appelle de la part de Datatym. Vous avez récemment répondu à une enquête : je souhaite simplement vérifier qu'elle s'est bien déroulée. Cela prend moins de deux minutes."],
        ["2. Entretien", "Avez-vous bien été interrogé(e) par un enquêteur de Datatym ? À quelle date, à peu près ? Combien de temps a duré l'entretien ?"],
        ["3. Réponses", "Reposer les questions de la fiche et comparer avec les réponses enregistrées."],
        ["4. Résultat", "Confirmé : tout concorde. Confirmé avec écarts : l'entretien a eu lieu mais des réponses diffèrent. "
                        "Entretien non réalisé : le répondant n'a pas été interrogé (le questionnaire sera retiré de la base)."],
        ["5. Clôture", "Remercier. Ne jamais redemander les réponses du questionnaire au-delà des questions de la fiche."],
    ], {"Étape": 16, "Ce qu'on dit ou fait": 110})
    wb.properties.creator = wb.properties.lastModifiedBy = auteur_fichiers()
    wb.save(chemin)
    return chemin
