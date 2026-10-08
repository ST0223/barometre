"""Sorties tabulaires du questionnaire : dictionnaire des variables, XLSForm Kobo, liste des enquêteurs, import QuestionPro."""
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .banque import NSP_CODE, TYPES, nom_barometre, nom_court, slug, sont_des_pays
from .config import MARINE

ENTETE_FILL = PatternFill("solid", fgColor=MARINE)
ENTETE_FONT = Font(bold=True, color="FFFFFF", name="Calibri")
OBLIGATOIRE = "Réponse obligatoire."
FICHIER_ENQUETEURS = "liste_enqueteur.csv"
EXCLUSIVES = (NSP_CODE, "aucun", "aucune", "non")


def _feuille(ws, colonnes, lignes, largeurs=None):
    ws.append(colonnes)
    for c in ws[1]:
        c.fill, c.font = ENTETE_FILL, ENTETE_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for l in lignes:
        ws.append(l)
    for i, col in enumerate(colonnes, 1):
        w = (largeurs or {}).get(col) or min(60, max(10, max((len(str(r[i - 1])) for r in lignes if i - 1 < len(r)), default=10) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


# ------------------------------------------------------------------ éléments partagés
def _indicateurs(cadrage):
    return {i["id"]: i for i in (cadrage or {}).get("indicateurs", []) if i.get("actif", True)}


def _q_par_id(quest):
    return {q["id"]: q for q in quest["questions_plates"]}


def _calculs_indicateurs(quest, cadrage):
    """Variables calculées ind_… : 1 si la réponse est favorable à l'indicateur, sinon 0."""
    inds = _indicateurs(cadrage)
    calculs = []
    for q in quest["questions_plates"]:
        ind = inds.get(q.get("indicateur"))
        if not ind or not TYPES[ind["type"]].get("positifs"):
            continue
        condition = " or ".join(f"${{{q['var']}}} = '{c}'" for c in TYPES[ind["type"]]["positifs"])
        calculs.append((f"ind_{ind['id'].lower()}", f"if({condition}, 1, 0)",
                        f"1 si réponse favorable à « {ind['libelle']} » ({q['num']}), sinon 0"))
    return calculs


def liste_enqueteurs(chemin, quest, enqueteurs=None):
    """Fichier média appelé par le XLSForm (select_one_from_file, pulldata) : name, label, superviseur, categorie."""
    if enqueteurs:
        lignes = [(e["code"], e["nom"], e.get("superviseur", ""), e.get("categorie", "principal")) for e in enqueteurs]
    else:
        q = _q_par_id(quest).get("id_enqueteur")
        lignes = [(m["code"], m["label"], "", "principal") for m in (q or {}).get("modalites", [])]
    with open(chemin, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["name", "label", "superviseur", "categorie"])
        w.writerows(lignes)
    return chemin


# ------------------------------------------------------------------ dictionnaire
def dictionnaire(chemin, quest, cadrage):
    inds = {i["id"]: i for i in cadrage.get("indicateurs", [])}
    lib = {pr["code"]: pr["libelle"] for pr in quest["profils"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Variables"
    lignes = [["consentement", "Acceptez-vous de participer à cette étude ?", "select_one", "oui = Oui | non = Non", "", "Accueil", "", "Tous"]]
    for q in quest["questions_plates"]:
        mods = " | ".join(f"{m['code']} = {m['label']}" for m in q.get("modalites", []))
        ind = q.get("indicateur") or ""
        lignes.append([q["var"], q["texte"], q["type"], mods, f"{ind} {inds[ind]['libelle']}" if ind in inds else "",
                       q["section"], q.get("condition", ""), ", ".join(lib.get(c, c) for c in q["profils"]) if q.get("profils") else "Tous"])
        if q["id"] == "id_enqueteur":
            for suffixe, sens in (("nom", "Nom de l'enquêteur"), ("sup", "Superviseur de l'enquêteur"), ("cat", "Catégorie de l'enquêteur")):
                lignes.append([f"{q['var']}_{suffixe}", f"{sens} (lu dans {FICHIER_ENQUETEURS})", "calculate", "", "", q["section"], "", ""])
        if q["type"] == "select_multiple":
            for m in q["modalites"]:
                lignes.append([f"{q['var']}_{m['code']}", f"{q['texte']} — {m['label']}", "binaire (0/1)", "0 = Non cité | 1 = Cité", ind, q["section"], q.get("condition", ""), ""])
            if any(m["code"] == "autre" for m in q["modalites"]):
                lignes.append([f"{q['var']}_autre", f"{q['num']} — Préciser autre", "text", "", ind, q["section"], f"Si Autre coché en {q['num']}", ""])
    for v, l in (("enqueteur", "Code de l'enquêteur (copie de la question d'identification)"), ("start", "Début de l'entretien"),
                 ("end", "Fin de l'entretien"), ("today", "Date de l'entretien"),
                 ("duree_minutes", "Durée de l'entretien en minutes (calculée)"), ("_submission_time", "Date de soumission")):
        lignes.append([v, l, "automatique", "", "", "Système", "", ""])
    for nom, calcul, sens in _calculs_indicateurs(quest, cadrage):
        lignes.append([nom, sens, "calculate", "0 = Non | 1 = Oui", nom.replace("ind_", "").upper(), "Calculs", calcul, ""])
    _feuille(ws, ["name", "label", "type", "modalités (code = libellé)", "indicateur rattaché", "section", "condition d'affichage", "profils concernés"],
             lignes, {"label": 60, "modalités (code = libellé)": 55, "condition d'affichage": 30})

    ws2 = wb.create_sheet("Indicateurs")
    l2 = []
    for i in cadrage.get("indicateurs", []):
        if not i.get("actif", True):
            continue
        vars_ = ", ".join(q["var"] for q in quest["questions_plates"] if q.get("indicateur") == i["id"])
        if i["type"] == "indice":
            vars_ = "Composantes : " + ", ".join(i.get("composantes", []))
        l2.append([i["id"], i["libelle"], i.get("definition", ""), TYPES[i["type"]]["libelle"], TYPES[i["type"]]["calcul"], vars_])
    _feuille(ws2, ["code", "indicateur", "définition", "type de mesure", "mode de calcul", "variables"], l2,
             {"définition": 50, "mode de calcul": 70, "indicateur": 40})
    wb.properties.creator = wb.properties.lastModifiedBy = "DATATYM"
    wb.save(chemin)
    return chemin


# ------------------------------------------------------------------ XLSForm KoboToolbox
COLONNES_SURVEY = ["type", "name", "label", "hint", "required", "required_message", "relevant", "constraint",
                   "constraint_message", "appearance", "choice_filter", "calculation"]


def _apparence(q):
    """columns-3 pour les listes courtes et nombreuses, columns-2 sinon, rien pour les libellés longs."""
    if q["type"] == "text":
        return "multiline" if q["id"] == "cl_commentaire" or not q.get("fixe") else ""
    if q["type"] not in ("select_one", "select_multiple") or not q.get("modalites"):
        return ""
    plus_long = max(len(m["label"]) for m in q["modalites"])
    if plus_long <= 16 and len(q["modalites"]) >= 4:
        return "columns-3"
    return "columns-2" if plus_long <= 40 else ""


def _noms_listes(quest, cadrage):
    """Nom de la liste de choix de chaque question ; les listes identiques sont partagées (oui_non, echelle_jugement…)."""
    inds = {i["id"]: i for i in (cadrage or {}).get("indicateurs", [])}
    fixes = {"id_zone": "pays" if sont_des_pays([m["label"] for m in _q_par_id(quest).get("id_zone", {}).get("modalites", [])]) else "zone",
             "id_profil": "fonction", "id_secteur": "secteur", "id_anciennete": "anciennete", "id_effectif": "effectif",
             "cl_resultats": "oui_non"}
    bases = {"jugement": "echelle_jugement", "satisfaction": "echelle_satisfaction", "intention": "intention",
             "evolution": "evolution", "perspective": "evolution"}
    par_signature = {(("oui", "Oui"), ("non", "Non")): "oui_non"}  # liste du consentement, réutilisée si identique
    noms, utilises = {}, {"oui_non"}
    for q in quest["questions_plates"]:
        if q["type"] not in ("select_one", "select_multiple") or q["id"] == "id_enqueteur":
            continue
        signature = tuple((m["code"], m["label"]) for m in q["modalites"])
        if signature in par_signature:
            noms[q["id"]] = par_signature[signature]
            continue
        ind = inds.get(q.get("indicateur"))
        base = fixes.get(q["id"]) or (bases.get(ind["type"]) if ind else None) or slug(ind["libelle"] if ind else q["texte"], maxi=24)
        nom, k = base, 2
        while nom in utilises:
            nom, k = f"{base}_{k}", k + 1
        utilises.add(nom)
        par_signature[signature] = noms[q["id"]] = nom
    return noms, par_signature


def xlsform(chemin, projet, quest, cadrage=None, enqueteurs=None):
    """Formulaire XLSForm complet pour KoboToolbox, sur le modèle du formulaire de référence Datatym.

    Accueil et consentement, groupe grp_consent (si consentement = oui) puis grp_eligible (profil dans la cible),
    sections en groupes, enquêteur lu dans liste_enqueteur.csv (select_one_from_file + pulldata), libellés en gras,
    messages d'obligation, apparences, exclusivités, variables calculées d'indicateurs.
    """
    fiche = projet["fiche"]
    par_id = _q_par_id(quest)
    noms, par_signature = _noms_listes(quest, cadrage)
    survey = []

    def ligne(**kw):
        survey.append([kw.get(c, "") for c in COLONNES_SURVEY])

    def question(q):
        var = q["var"]
        obligatoire = q.get("obligatoire", True)
        req = {"required": "yes" if obligatoire else "", "required_message": OBLIGATOIRE if obligatoire else ""}
        label = f"**{q['num']}.** {q['texte']}"
        if q["id"] == "id_enqueteur":
            ligne(type=f"select_one_from_file {FICHIER_ENQUETEURS}", name=var, label=label, appearance="minimal",
                  hint=q.get("consigne", "").replace(" Liste mise à jour à l'étape échantillonnage.", ""), **req)
            for suffixe, colonne in (("nom", "label"), ("sup", "superviseur"), ("cat", "categorie")):
                ligne(type="calculate", name=f"{var}_{suffixe}",
                      calculation=f"pulldata('{FICHIER_ENQUETEURS[:-4]}', '{colonne}', 'name', ${{{var}}})")
            return
        contraintes, messages = [], []
        if q["type"] in ("select_one", "select_multiple"):
            type_ = f"{q['type']} {noms[q['id']]}"
        elif q["type"] == "email":
            type_ = "text"
            contraintes.append("regex(., '^[^@\\s]+@[^@\\s]+\\.[^@\\s]+$')")
            messages.append("Adresse e-mail invalide.")
        else:
            type_ = q["type"]
        if q["type"] == "select_multiple":
            if q.get("max_choix"):
                contraintes.append(f"count-selected(.) <= {q['max_choix']}")
                messages.append(f"{q['max_choix']} réponses au maximum.")
            for m in q["modalites"]:
                if m["code"] in EXCLUSIVES:
                    contraintes.append(f"not(selected(., '{m['code']}') and count-selected(.) > 1)")
                    messages.append(f"« {m['label']} » est exclusive des autres réponses.")
        ligne(type=type_, name=var, label=label, relevant=q.get("relevant", ""), constraint=" and ".join(contraintes),
              constraint_message=" ".join(messages), appearance=_apparence(q), **req)
        if q["type"] == "select_multiple" and any(m["code"] == "autre" for m in q["modalites"]):
            ligne(type="text", name=f"{var}_autre", label="Préciser Autre", required="yes", required_message=OBLIGATOIRE,
                  relevant=f"selected(${{{var}}}, 'autre')")
        if q["id"] == "id_profil":
            ligne(type="note", name="hors_cible", label="**Merci.** Ce baromètre ne concerne pas votre fonction.",
                  relevant=f"${{{var}}} = 'autre_hors'")

    ligne(type="start", name="start")
    ligne(type="end", name="end")
    ligne(type="today", name="today")
    ligne(type="begin_group", name="accueil")
    ligne(type="note", name="message", label=f"**{nom_barometre(fiche['theme'])}**\n\n" + quest["consentement"].replace("Datatym", "**Datatym**", 1))
    ligne(type="select_one oui_non", name="consentement", label="Acceptez-vous de participer à cette étude ?", required="yes",
          required_message=OBLIGATOIRE, appearance="columns-2")
    ligne(type="note", name="consentement_fin", label="**Merci de votre attention.**", relevant="${consentement} = 'non'")
    ligne(type="end_group")
    ligne(type="begin_group", name="grp_consent", relevant="${consentement} = 'oui'")
    q_profil = par_id.get("id_profil")
    eligibilite_ouverte = False
    for s in quest["sections"]:
        if s["lettre"] != "A" and q_profil and not eligibilite_ouverte:
            ligne(type="begin_group", name="grp_eligible", relevant=f"${{{q_profil['var']}}} != 'autre_hors'")
            eligibilite_ouverte = True
        ligne(type="begin_group", name=f"sec_{s['lettre'].lower()}", label=f"**Section {s['lettre']} : {s['titre']}**")
        for q in s["ids"]:
            question(q)
        if s is quest["sections"][-1]:
            ligne(type="note", name="merci", label="**Merci pour votre participation.**")
        ligne(type="end_group")
    var_enq = par_id.get("id_enqueteur", {}).get("var")
    if var_enq:
        ligne(type="calculate", name="enqueteur", calculation=f"${{{var_enq}}}")
    ligne(type="calculate", name="duree_minutes", calculation="round((decimal-date-time(${end}) - decimal-date-time(${start})) * 1440, 1)")
    for nom, calcul, _ in _calculs_indicateurs(quest, cadrage):
        ligne(type="calculate", name=nom, calculation=calcul)
    if eligibilite_ouverte:
        ligne(type="end_group")
    ligne(type="end_group")

    choices = [[liste, code, label] for signature, liste in par_signature.items() for code, label in signature]
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    _feuille(ws, COLONNES_SURVEY, survey, {"label": 60, "relevant": 40, "constraint": 40, "calculation": 50, "hint": 30})
    _feuille(wb.create_sheet("choices"), ["list_name", "name", "label"], choices, {"label": 45})
    _feuille(wb.create_sheet("settings"), ["form_title", "form_id", "version", "style", "allow_choice_duplicates"],
             [[nom_barometre(fiche["theme"]), nom_court(fiche["theme"]), datetime.now().strftime("%Y%m%d%H%M"), "pages", "yes"]])
    wb.properties.creator = wb.properties.lastModifiedBy = "DATATYM"
    wb.save(chemin)
    return chemin


# ------------------------------------------------------------------ QuestionPro
def questionpro(chemin, quest):
    """Fichier texte pour l'import en masse de QuestionPro : une question par paragraphe, une réponse par ligne."""
    lignes = [quest["consentement"], "", "Acceptez-vous de participer à cette étude ?", "Oui", "Non", ""]
    for s in quest["sections"]:
        for q in s["ids"]:
            if q.get("terrain_seul"):
                continue
            entete = f"{q['num']}. {q['texte']}"
            if q["type"] == "select_multiple":
                entete += f" (plusieurs réponses possibles, {q.get('max_choix', 3)} maximum)"
            lignes.append(entete)
            for m in q.get("modalites", []):
                lignes.append(m["label"])
            lignes.append("")
    with open(chemin, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lignes))
    return chemin
