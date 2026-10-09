"""PowerPoint de formation des enquêteurs (.pptx modifiable), généré pour chaque baromètre.

Construit sur le masque de l'exemplaire Datatym (datatym/modeles/formation/modele.pptx : dispositions
COUVERTURE, INTERCALAIRE et CONTENU, logos Datatym et Agiltym, pied de page) et sur ses icônes.
"""
import math
from copy import deepcopy
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from .banque import nom_barometre, sont_des_pays
from .etapes.questionnaire import consignes_terrain
from .outils import MOIS, date_fr

MODELE = Path(__file__).resolve().parent / "modeles" / "formation"
MARINE, BLEU, OR, CLAIR, JAUNE = "19243E", "0087E0", "E0B61F", "EEF4FB", "FCF6DE"
GRIS, TEXTE, BRUN, CONSIGNE, BLEU_PALE = "6B7280", "1F2937", "8A6A00", "B4381F", "C9D6EA"
TYPE_LIB = {"select_one": "Choix unique", "select_multiple": "Choix multiple", "text": "Texte libre",
            "integer": "Nombre", "email": "E-mail", "note": "Message"}


def _po(x):
    """Pouces → EMU."""
    return Emu(int(x * 914400))


def _c(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, pied):
        self.prs = Presentation(MODELE / "modele.pptx")
        self.dispo = {l.name: l for l in self.prs.slide_master.slide_layouts}
        for sh in self.dispo["CONTENU"].shapes:  # pied de page propre à l'étude
            if sh.has_text_frame and "{{pied}}" in sh.text_frame.text:
                sh.text_frame.paragraphs[0].runs[0].text = pied
                sh.left = _po(2.4)  # un peu d'air après le logo Agiltym
        self.num = next(sh for sh in self.dispo["CONTENU"].shapes if sh.is_placeholder and sh.placeholder_format.type == 13)

    # ---------- diapositives ----------
    def couverture(self, titre, sous_titre, mention):
        s = self.prs.slides.add_slide(self.dispo["COUVERTURE"])
        s.placeholders[102].text = titre
        s.placeholders[103].text = sous_titre
        self.texte(s, 0.7, 6.55, 8.0, 0.35, mention, 14, couleur=BLEU_PALE)
        self.pastilles(s, 9.16, 3.0, 1.0)
        return s

    def intercalaire(self, num, titre, sous_titre):
        s = self.prs.slides.add_slide(self.dispo["INTERCALAIRE"])
        self.texte(s, 0.7, 1.55, 3.0, 1.3, f"{num:02d}", 72, True, OR)
        s.placeholders[102].text = titre
        s.placeholders[103].text = sous_titre
        self.pastilles(s, 8.72, 2.6, 1.08)
        return s

    def contenu(self, titre):
        s = self.prs.slides.add_slide(self.dispo["CONTENU"])
        s.placeholders[101].text = titre
        s.shapes._spTree.append(deepcopy(self.num._element))  # numéro de page du masque
        return s

    # ---------- briques ----------
    def forme(self, s, x, y, w, h, couleur, type_=MSO_SHAPE.ROUNDED_RECTANGLE, arrondi=0.08):
        f = s.shapes.add_shape(type_, _po(x), _po(y), _po(w), _po(h))
        f.fill.solid()
        f.fill.fore_color.rgb = _c(couleur)
        f.line.fill.background()
        f.shadow.inherit = False
        if type_ == MSO_SHAPE.ROUNDED_RECTANGLE:
            f.adjustments[0] = arrondi
        return f

    def texte(self, s, x, y, w, h, contenu, taille=14, gras=False, couleur=TEXTE, align=PP_ALIGN.LEFT,
              ancre=MSO_ANCHOR.TOP, italique=False, interligne=None):
        """contenu : texte, ou liste de paragraphes ; un paragraphe peut être une liste de runs (texte, options)."""
        tb = s.shapes.add_textbox(_po(x), _po(y), _po(w), _po(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = ancre
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
        paragraphes = contenu if isinstance(contenu, list) else [contenu]
        for i, para in enumerate(paragraphes):
            pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pp.alignment = align
            if interligne:
                pp.space_after = Pt(interligne)
            runs = para if isinstance(para, list) else [(para, {})]
            for txt, opt in runs:
                r = pp.add_run()
                r.text = str(txt)
                r.font.size = Pt(opt.get("taille", taille))
                r.font.bold = opt.get("gras", gras)
                r.font.italic = opt.get("italique", italique)
                r.font.color.rgb = _c(opt.get("couleur", couleur))
        return tb

    def pastilles(self, s, x0, y0, k):
        """Motif des pastilles du logo Datatym (couvertures et intercalaires)."""
        for dx, dy, w, rond in ((0.44, 0, 1.76, False), (2.70, 0, 0.46, True), (0, 0.86, 1.65, False),
                                (1.98, 0.86, 1.65, False), (0.82, 1.72, 2.20, False)):
            if rond:
                self.forme(s, x0 + dx * k, y0 + dy * k, 0.46 * k, 0.46 * k, "FFFFFF", MSO_SHAPE.OVAL)
            else:
                self.forme(s, x0 + dx * k, y0 + dy * k, w * k, 0.46 * k, BLEU, arrondi=0.5)

    def icone(self, s, x, y, nom, taille=0.85, fond=BLEU, rond=True):
        """Icône blanche dans un disque (ou icône seule si fond=None)."""
        if fond:
            self.forme(s, x, y, taille, taille, fond, MSO_SHAPE.OVAL if rond else MSO_SHAPE.ROUNDED_RECTANGLE)
        marge = taille * 0.25
        s.shapes.add_picture(str(MODELE / "icones" / f"{nom}.png"), _po(x + marge), _po(y + marge), _po(taille - 2 * marge))

    def soustitre(self, s, texte):
        self.texte(s, 0.6, 1.2, 12.0, 0.5, texte, 16, couleur=GRIS)

    def tableau(self, s, x, y, w, colonnes, lignes, largeurs=None, taille=12, total=False):
        h = _po(0.36 * (len(lignes) + 1))
        t = s.shapes.add_table(len(lignes) + 1, len(colonnes), _po(x), _po(y), _po(w), h).table
        if largeurs:
            tot = sum(largeurs)
            for j, l in enumerate(largeurs):
                t.columns[j].width = _po(w * l / tot)
        for i, ligne in enumerate([colonnes] + lignes):
            for j, val in enumerate(ligne):
                cell = t.cell(i, j)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _c(MARINE if i == 0 else (CLAIR if i % 2 == 0 else "FFFFFF"))
                cell.margin_top = cell.margin_bottom = Emu(45720)
                for par in cell.text_frame.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(taille)
                        r.font.bold = i == 0 or (total and i == len(lignes))
                        r.font.color.rgb = _c("FFFFFF" if i == 0 else TEXTE)
        return t


def _icone_profil(libelle):
    l = libelle.lower()
    for mots, nom in ((("dsi", "digital", "informat", "cdo", "tech"), "informatique"),
                      (("commerc", "vendeur", "marketing", "vente"), "commercial"),
                      (("daf", "financ", "dfc", "compta"), "finance"),
                      (("drh", "rh", "ressources humaines", "salari", "employ", "manager", "jeune", "ménage"), "equipe"),
                      (("dg", "directeur", "dirigeant", "pdg", "responsable", "gérant", "chef"), "directeur")):
        if any(m in l for m in mots):
            return nom
    return "badge"


def generer_pptx(chemin, projet, cadrage, quest, ech):
    fiche = projet["fiche"]
    nom = nom_barometre(fiche["theme"])
    d = Deck(f"{nom} · Formation des enquêteurs")
    kobo = fiche.get("outil", "kobo") == "kobo"
    outil = "KoboToolbox (KoboCollect)" if kobo else "QuestionPro"
    jours = [date.fromisoformat(j) for j in ech["jours"]]
    profils = quest["profils"]
    lib = {pr["code"]: pr["libelle"] for pr in profils}
    toutes = quest["questions_plates"]
    terrain = toutes
    digital = [q for q in toutes if not q.get("terrain_seul")]
    consignes = consignes_terrain(quest)
    pays = sont_des_pays(ech["zones"])
    zone_mot = "pays" if pays else "zones"
    duree = math.ceil(quest.get("duree_max") or 0) or 12
    mois = f"{MOIS[date.today().month - 1].capitalize()} {date.today().year}"
    derniere = quest["sections"][-1]["lettre"]

    # ------------------------------------------------ couverture et programme
    d.couverture(nom, "Formation des enquêteurs · Questionnaire versions terrain et digitale",
                 f"DATATYM, marque d'AGILTYM  ·  {mois}")
    s = d.contenu("Programme de la formation")
    modules = [("L'étude", "Objectif, cible, chiffres clés"), ("Qui interroger", "Profils et éligibilité du répondant"),
               ("Le questionnaire", "Deux versions, règles de passation"), ("Section par section", f"Les questions de A à {derniere}"),
               ("L'outil et les quotas", f"{'KoboToolbox' if kobo else 'QuestionPro'}, objectifs, contrôle qualité"),
               ("Points de vigilance", "Sauts, réponses exclusives, mise en pratique")]
    for i, (t, st) in enumerate(modules):
        x, y = 0.6 + (i % 3) * 4.1, 1.55 + (i // 3) * 2.6
        d.forme(s, x, y, 3.8, 2.25, CLAIR)
        d.texte(s, x + 0.3, y + 0.25, 1.2, 0.7, f"{i + 1:02d}", 36, True, BLEU)
        d.texte(s, x + 0.3, y + 1.0, 3.3, 0.45, t, 20, True, MARINE)
        d.texte(s, x + 0.3, y + 1.45, 3.3, 0.6, st, 14, couleur=GRIS)

    # ------------------------------------------------ 01 L'étude
    d.intercalaire(1, "L'étude et sa cible", "Pourquoi ce baromètre, qui interroger")
    s = d.contenu("L'étude en chiffres")
    d.soustitre(s, fiche["objectif"])
    chiffres = [(str(fiche["taille"]), "questionnaires à réaliser"), (str(len(ech["zones"])), f"{zone_mot} d'enquête"),
                (str(len(jours)), "jours de collecte"), (f"{duree} min", "durée approximative du questionnaire")]
    for i, (v, l) in enumerate(chiffres):
        x = 0.6 + i * 3.08
        d.forme(s, x, 2.1, 2.8, 2.3, MARINE)
        d.texte(s, x + 0.25, 2.35, 2.3, 1.0, v, 48, True, OR)
        d.texte(s, x + 0.25, 3.4, 2.3, 0.8, l, 15, couleur="FFFFFF")
    d.forme(s, 0.6, 4.8, 12.1, 1.6, JAUNE)
    d.texte(s, 0.9, 4.98, 3.0, 0.35, zone_mot.capitalize(), 14, True, MARINE)
    d.texte(s, 0.9, 5.32, 11.5, 0.45, " · ".join(ech["zones"]), 14, couleur=TEXTE)
    d.texte(s, 0.9, 5.85, 11.5, 0.4, "Les réponses sont confidentielles et ne sont jamais publiées de façon nominative.", 14, True, BRUN)

    s = d.contenu("Qui interroger ?")
    d.soustitre(s, "Un répondant = une personne du profil attendu, dans la zone attribuée ; un seul questionnaire par personne.")
    n = max(len(profils), 1)
    larg = min(2.25, (12.1 - 0.21 * (n - 1)) / n)
    for i, pr in enumerate(profils[:5]):
        x = 0.6 + i * (larg + 0.21)
        d.forme(s, x, 1.95, larg, 2.2, CLAIR)
        d.icone(s, x + larg / 2 - 0.425, 2.15, _icone_profil(pr["libelle"]))
        d.texte(s, x + 0.15, 3.15, larg - 0.3, 0.85,
                [pr["libelle"], [("Direction" if pr["niveau"] == "direction" else "Opérationnel", {"taille": 12, "gras": False, "couleur": GRIS})]],
                14, True, MARINE, PP_ALIGN.CENTER)
    q_profil = next((q for q in toutes if q["id"] == "id_profil"), None)
    blocs = [("Vérifier le profil d'abord", f"Question {q_profil['var'] if q_profil else 'a3'} : si « Autre fonction », remercier et terminer."),
             (f"Dans le bon {('pays' if pays else 'zone')}", "Chaque enquêteur travaille dans la zone qui lui est attribuée."),
             ("Secteur visé", f"{fiche.get('secteur') or 'Secteur formel'} ; un seul répondant par structure.")]
    for i, (t, txt) in enumerate(blocs):
        d.texte(s, 0.6 + i * 4.1, 4.55, 3.8, 0.4, t, 16, True, MARINE)
        d.texte(s, 0.6 + i * 4.1, 5.0, 3.8, 1.0, txt, 14)

    # ------------------------------------------------ 02 Le questionnaire
    d.intercalaire(2, "Le questionnaire", "Deux versions, des règles de passation communes")
    s = d.contenu("Deux versions du même questionnaire")
    for i, (titre, nb, puces, fond, ct) in enumerate((
            ("Version terrain", len(terrain), ["L'enquêteur lit chaque question et ses réponses", "Toutes les questions, y compris la question « Enquêteur »",
                                                f"Saisie sur {outil}, hors ligne possible"], MARINE, "FFFFFF"),
            ("Version digitale", len(digital), ["Le répondant remplit seul, par lien ou code QR", "Mêmes codes, libellés et sauts",
                                                 "L'enquêteur est reconnu par son lien personnel"], CLAIR, MARINE))):
        x = 0.6 + i * 6.15
        d.forme(s, x, 1.45, 5.95, 3.9, fond)
        d.texte(s, x + 0.35, 1.7, 5.0, 0.5, titre, 20, True, ct)
        d.texte(s, x + 0.35, 2.25, 5.0, 0.8, [[(str(nb), {"taille": 40, "gras": True, "couleur": OR if i == 0 else BLEU}),
                                                (" questions", {"taille": 16, "couleur": ct})]], 16)
        d.texte(s, x + 0.35, 3.2, 5.3, 2.3, [f"•  {p}" for p in puces], 15, couleur=ct, interligne=8)
    d.texte(s, 0.6, 5.65, 12.1, 0.4, f"Durée estimée : environ {duree} minutes en version terrain. Les deux versions alimentent la même base.", 13,
            couleur=GRIS, italique=True)

    s = d.contenu("Les règles de passation")
    regles = [("lire", "Lire la question en gras, puis les réponses", "Lire le texte tel qu'il est écrit, sans le reformuler."),
              ("ne_pas_lire", "Ne jamais lire la consigne", "La consigne indique le type de question, les sauts et les limites. Elle ne se lit pas."),
              ("obligatoire", "Toutes les questions sont obligatoires", "Sauf celles marquées « Facultatif »."),
              ("chronometre", f"Environ {duree} minutes", "Annoncer la durée au démarrage et la respecter.")]
    for i, (ic, t, txt) in enumerate(regles):
        x, y = 0.6 + (i % 2) * 6.2, 1.55 + (i // 2) * 2.45
        d.icone(s, x, y, ic, 0.9)
        d.texte(s, x + 1.15, y - 0.05, 4.75, 0.75, t, 18, True, MARINE)
        d.texte(s, x + 1.15, y + 0.75, 4.75, 1.2, txt, 15)

    types = [t for t in ("select_one", "select_multiple", "text", "integer", "email") if any(q["type"] == t for q in toutes)]
    desc = {"select_one": ("choix_unique", "Une seule réponse possible. Lire les modalités, sauf « Ne sait pas »."),
            "select_multiple": ("choix_multiple", "Plusieurs réponses, dans la limite indiquée (3 au maximum)."),
            "text": ("texte_libre", "Noter la réponse mot pour mot, sans la reformuler."),
            "integer": ("classement", "Un nombre ; vérifier l'ordre de grandeur avec le répondant."),
            "email": ("texte_libre", "Adresse e-mail professionnelle ; vérifier l'orthographe.")}
    s = d.contenu(f"Les {len(types)} types de questions")
    for i, t in enumerate(types):
        y = 1.4 + i * 0.95
        d.forme(s, 0.6, y, 12.1, 0.8, CLAIR if i % 2 == 0 else "FFFFFF")
        d.icone(s, 0.75, y + 0.12, desc[t][0], 0.56, fond=None)
        d.texte(s, 1.55, y, 3.0, 0.8, TYPE_LIB[t], 15, True, MARINE, ancre=MSO_ANCHOR.MIDDLE)
        d.texte(s, 4.6, y, 8.0, 0.8, desc[t][1], 14, ancre=MSO_ANCHOR.MIDDLE)

    s = d.contenu(f"Le questionnaire en {len(quest['sections'])} sections")
    lignes = [[sec["lettre"], sec["titre"], len(sec["ids"]), len([q for q in sec["ids"] if not q.get("terrain_seul")])] for sec in quest["sections"]]
    lignes.append(["", "Total", len(terrain), len(digital)])
    d.tableau(s, 0.6, 1.4, 12.1, ["Section", "Thème", "Version terrain", "Version digitale"], lignes, [0.8, 6, 1.6, 1.6], 13, total=True)

    # ------------------------------------------------ 03 Section par section
    d.intercalaire(3, "Section par section", f"Les questions de A à {derniere}, avec leurs consignes")
    for sec in quest["sections"]:
        qs = sec["ids"]
        for k in range(0, len(qs), 6):
            s = d.contenu(f"Section {sec['lettre']} · {sec['titre']}" + (" (suite)" if k else ""))
            for j, q in enumerate(qs[k:k + 6]):
                x, y = 0.6 + (j % 2) * 6.15, 1.25 + (j // 2) * 1.7
                reservee = bool(q.get("profils"))
                d.forme(s, x, y, 5.95, 1.55, JAUNE if reservee else CLAIR)
                d.forme(s, x + 0.15, y + 0.12, 0.75, 0.32, BLEU, arrondi=0.5)
                d.texte(s, x + 0.15, y + 0.12, 0.75, 0.32, q["var"], 13, True, "FFFFFF", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                d.texte(s, x + 1.05, y + 0.12, 2.6, 0.32, TYPE_LIB.get(q["type"], q["type"]), 12, couleur=GRIS, ancre=MSO_ANCHOR.MIDDLE)
                if reservee:
                    etiquette = ", ".join(lib.get(c, c) for c in q["profils"])
                    d.forme(s, x + 3.75, y + 0.12, 2.05, 0.32, OR, arrondi=0.5)
                    d.texte(s, x + 3.75, y + 0.12, 2.05, 0.32, etiquette[:24], 11, True, MARINE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                cons = [c for c in consignes[q["id"]][1:] if not c.startswith("Poser seulement si") or not reservee]
                taille = 14 if len(q["texte"]) < 110 else 12
                d.texte(s, x + 0.15, y + 0.52, 5.65, 0.97,
                        [q["texte"]] + ([[(" ".join(cons)[:150], {"taille": 11, "gras": False, "italique": True, "couleur": CONSIGNE})]] if cons else []),
                        taille, True, MARINE)

    # ------------------------------------------------ 04 L'outil et les quotas
    d.intercalaire(4, "L'outil et les quotas", f"{outil}, objectifs par enquêteur, contrôle qualité")
    etapes = ([("Installer", "Télécharger KoboCollect sur le Play Store et autoriser l'application."),
               ("Configurer", "Paramètres › Serveur : saisir l'adresse et l'identifiant fournis par Datatym."),
               ("Télécharger", "« Télécharger formulaire » › cocher le baromètre ; à refaire à chaque nouvelle version."),
               ("Saisir", "« Remplir formulaire » : suivre les écrans, les sauts sont automatiques."),
               ("Envoyer", "« Envoyer formulaire finalisé » chaque soir, ou dès qu'une connexion est disponible."),
               ("Hors ligne", "La saisie fonctionne sans internet ; les formulaires restent sur le téléphone jusqu'à l'envoi.")] if kobo else
              [("Installer", "Installer l'application QuestionPro hors ligne sur la tablette ou le smartphone."),
               ("Se connecter", "Utiliser l'identifiant enquêteur fourni par Datatym."),
               ("Synchroniser", "Récupérer le questionnaire du baromètre."),
               ("Saisir", "Ouvrir l'enquête et suivre les écrans ; les sauts sont automatiques."),
               ("Envoyer", "Synchroniser chaque soir pour transmettre les réponses."),
               ("Hors ligne", "La saisie fonctionne sans internet ; synchroniser dès qu'une connexion est disponible.")])
    s = d.contenu("Prise en main de l'outil en 6 étapes")
    for i, (t, txt) in enumerate(etapes):
        x, y = 0.6 + (i % 3) * 4.1, 1.45 + (i // 3) * 2.6
        d.forme(s, x, y, 3.8, 2.35, CLAIR)
        d.texte(s, x + 0.3, y + 0.22, 3.3, 0.5, [[(f"{i + 1:02d}  ", {"couleur": BLEU, "taille": 22}), (t, {"couleur": MARINE})]], 18, True)
        d.texte(s, x + 0.3, y + 0.9, 3.25, 1.3, txt, 14)

    codes = [pr["code"] for pr in profils]
    for k in range(0, len(ech["enqueteurs"]), 12):
        s = d.contenu("Objectifs par enquêteur" + (" (suite)" if k else ""))
        lignes = [[m["nom"], m["zone"]] + [m["par_profil"][c] for c in codes] + [m["objectif"], str(round(m["objectif"] / len(jours), 1)).replace(".", ",")]
                  for m in ech["enqueteurs"][k:k + 12]]
        d.tableau(s, 0.6, 1.35, 12.1, ["Enquêteur", "Pays" if pays else "Zone"] + [lib[c] for c in codes] + ["Total", "Par jour"], lignes, taille=12)
    d.texte(s, 0.6, 6.3, 12.1, 0.4, f"Collecte du {date_fr(jours[0])} au {date_fr(jours[-1])}, formation la veille du démarrage.", 13, couleur=GRIS, italique=True)

    s = d.contenu("Suivi et contrôle qualité")
    controles = [("telephone", "Chaque soir", "Envoi des formulaires et point avec le superviseur : réalisé contre objectif."),
                 ("filtre", "Quota atteint", "On n'interroge plus ce profil dans cette zone ; priorité aux quotas restants."),
                 ("bouclier", "Contrôle", "10 % des répondants sont rappelés pour vérifier l'entretien."),
                 ("arret", "Rejet", "Un questionnaire trop rapide ou incohérent est annulé et doit être refait.")]
    for i, (ic, t, txt) in enumerate(controles):
        x, y = 0.6 + (i % 2) * 6.2, 1.55 + (i // 2) * 2.45
        d.icone(s, x, y, ic, 0.9)
        d.texte(s, x + 1.15, y - 0.05, 4.75, 0.75, t, 18, True, MARINE)
        d.texte(s, x + 1.15, y + 0.75, 4.75, 1.2, txt, 15)

    # ------------------------------------------------ 05 Points de vigilance
    d.intercalaire(5, "Points de vigilance", "Sauts, réponses exclusives et mise en pratique")
    sauts = [["Consentement", "« Non »", "Remercier et terminer"]]
    if q_profil:
        sauts.append([q_profil["var"], "« Autre fonction »", "Remercier et terminer (hors cible)"])
    for q in toutes:
        if q.get("profils"):
            sauts.append([q["var"], "Profil " + " / ".join(lib.get(c, c) for c in q["profils"]), f"Poser {q['var']} ; sinon passer"])
        if q.get("condition_sur"):
            parent = next((x for x in toutes if x["id"] == q["condition_sur"][0]), None)
            if parent:
                lab = next((m["label"] for m in parent["modalites"] if m["code"] == q["condition_sur"][1]), q["condition_sur"][1])
                sauts.append([parent["var"], f"« {lab} »", f"Poser {q['var']}"])
    s = d.contenu("Les sauts à connaître par cœur")
    d.tableau(s, 0.6, 1.35, 12.1, ["Question", "Si la réponse est…", "Alors"], sauts[:13], [1.2, 3.5, 3.5], 13)

    s = d.contenu("Réponses exclusives et nombre de réponses")
    d.forme(s, 0.6, 1.4, 6.2, 4.6, CLAIR)
    d.icone(s, 0.85, 1.6, "choix_unique", 0.6, fond=None)
    d.texte(s, 1.6, 1.62, 5.0, 0.5, "Réponses exclusives", 18, True, MARINE)
    excl = []
    for q in toutes:
        for m in q.get("modalites", []):
            if q["type"] == "select_multiple" and (m["code"] in ("99", "aucun", "aucune") or m["label"].lower().startswith(("aucun", "ne sait pas"))):
                excl.append([(f"{q['var']}  ", {"gras": True, "couleur": BLEU}), (f"« {m['label']} »", {})])
    d.texte(s, 0.95, 2.35, 5.6, 3.5, ["Dans une question à choix multiple, ces réponses ne se cochent jamais avec une autre :"]
            + (excl or [[("« Ne sait pas »", {"gras": True, "couleur": BLEU}), (" et « Aucun » : réponse seule, jamais combinée.", {})]])
            + [[("Plusieurs réponses cochées avec l'une d'elles : relire la question au répondant.", {"italique": True, "couleur": GRIS, "taille": 13})]],
            14, interligne=10)
    d.forme(s, 7.0, 1.4, 5.7, 4.6, MARINE)
    d.texte(s, 7.3, 1.62, 5.0, 0.5, "Nombre maximum de réponses", 18, True, "FFFFFF")
    multiples = [q for q in toutes if q.get("max_choix")]
    for i, q in enumerate(multiples[:4]):
        d.texte(s, 7.3, 2.35 + i * 0.85, 1.0, 0.7, str(q["max_choix"]), 30, True, OR)
        d.texte(s, 8.3, 2.42 + i * 0.85, 4.2, 0.7, [q["var"], [("réponses au maximum", {"taille": 12, "gras": False, "couleur": BLEU_PALE})]], 15, True, "FFFFFF")

    s = d.contenu("Les réflexes de l'enquêteur")
    reflexes = [("badge", "Se présenter", "Badge visible, présentation de Datatym et du baromètre."),
                ("citation", "Lire le consentement en entier", "Sans « Oui » explicite, l'entretien s'arrête."),
                ("filtre", "Vérifier l'éligibilité d'abord", "Fonction et zone : hors cible, on remercie et on termine."),
                ("bouclier", "Rassurer sur la confidentialité", "Aucune réponse n'est publiée de façon nominative."),
                ("texte_libre", "Noter mot pour mot", "Pour « Autre (préciser) » et les questions ouvertes."),
                ("courriel", "Envoyer chaque soir", "Les formulaires non envoyés n'apparaissent pas dans le suivi.")]
    for i, (ic, t, txt) in enumerate(reflexes):
        x, y = 0.6 + (i % 2) * 6.2, 1.4 + (i // 2) * 1.7
        if ic == "texte_libre":
            d.forme(s, x, y, 0.8, 0.8, BLEU, MSO_SHAPE.OVAL)
            d.icone(s, x + 0.08, y + 0.08, "texte_libre", 0.64, fond=None)
        else:
            d.icone(s, x, y, ic, 0.8)
        d.texte(s, x + 1.05, y, 4.9, 0.45, t, 17, True, MARINE)
        d.texte(s, x + 1.05, y + 0.48, 4.9, 0.9, txt, 14)

    s = d.contenu("Cas pratiques, jeu de rôle et quiz")
    d.forme(s, 0.6, 1.4, 6.0, 4.9, CLAIR)
    d.texte(s, 0.9, 1.6, 5.4, 0.5, "Cas pratiques", 18, True, MARINE)
    d.texte(s, 0.9, 2.2, 5.4, 4.0, [
        "1. Le répondant est pressé et propose de remplir « plus tard ».",
        "2. Il demande ce que vous entendez par une question.",
        "3. Votre quota d'un profil est atteint, une personne de ce profil veut répondre.",
        "4. Pas de réseau pendant deux jours : que faire des formulaires ?",
        "Jeu de rôle en binôme : passation complète chronométrée."], 14, interligne=8)
    d.forme(s, 6.8, 1.4, 5.9, 4.9, MARINE)
    d.texte(s, 7.1, 1.6, 5.3, 0.5, "Quiz de fin de formation", 18, True, "FFFFFF")
    d.texte(s, 7.1, 2.2, 5.3, 4.0, [
        "1. Que faire si le consentement est refusé ?",
        f"2. Et si {q_profil['var'] if q_profil else 'la fonction'} = « Autre fonction » ?",
        "3. Combien de réponses au maximum pour un choix multiple ?",
        "4. Quand lit-on « Ne sait pas » ?",
        f"5. Votre objectif quotidien moyen ? (≈ {str(round(fiche['taille'] / len(jours) / max(ech['nb_enqueteurs'], 1), 1)).replace('.', ',')})",
        "6. Quand envoyer les formulaires ?"], 14, couleur="FFFFFF", interligne=8)

    d.couverture("Merci et bonne collecte", "Place aux questions et à la mise en situation", f"DATATYM, marque d'AGILTYM  ·  {mois}")

    cp = d.prs.core_properties
    cp.author = cp.last_modified_by = "DATATYM"
    cp.title, cp.subject = nom, "Formation des enquêteurs"
    d.prs.save(chemin)
    return chemin
