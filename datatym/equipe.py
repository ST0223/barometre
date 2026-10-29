"""Équipe d'enquêteurs recrutée par les RH : fichier Excel à renseigner, lecture et contrôle de la liste retournée.

Les RH renseignent une ligne par enquêteur retenu ; l'identifiant (ENQML01, ENQCI01…) est attribué par la plateforme
à l'import, pays par pays, et sert ensuite dans Kobo, QuestionPro (custom1) et le suivi par collecteur.
"""
import csv
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from .banque import slug
from .formulaires import _feuille

# colonne, clé, consigne
COLONNES = [
    ("Nom", "nom", "Nom de famille, en majuscules"),
    ("Prénom", "prenom", "Prénom(s) usuel(s)"),
    ("Pays de résidence", "pays", "Pays où l'enquêteur travaillera, choisi dans la liste déroulante"),
    ("Ville", "ville", "Ville de résidence"),
    ("Téléphone", "telephone", "Numéro joignable sur le terrain (WhatsApp de préférence), précédé de l'indicatif du pays "
     "et saisi avec une apostrophe devant ('+221 771234567), sinon Excel le transforme en nombre et supprime le +"),
    ("Mail", "mail", "Adresse personnelle et active : contrat, ordre de mission et notes de suivi y seront envoyés"),
]

INDICATIFS = {
    "senegal": "221", "mauritanie": "222", "mali": "223", "guinee": "224", "cote_d_ivoire": "225", "burkina_faso": "226",
    "niger": "227", "togo": "228", "benin": "229", "ghana": "233", "nigeria": "234", "tchad": "235", "centrafrique": "236",
    "republique_centrafricaine": "236", "cameroun": "237", "guinee_equatoriale": "240", "gabon": "241", "congo": "242",
    "republique_du_congo": "242", "congo_brazzaville": "242", "rd_congo": "243", "rdc": "243",
    "republique_democratique_du_congo": "243", "guinee_bissau": "245", "maroc": "212", "tunisie": "216", "france": "33",
}
# code pays des identifiants (ENQML01, ENQCI01, ENQSN01…)
CODES_PAYS = {
    "senegal": "SN", "mauritanie": "MR", "mali": "ML", "guinee": "GN", "cote_d_ivoire": "CI", "burkina_faso": "BF",
    "niger": "NE", "togo": "TG", "benin": "BJ", "ghana": "GH", "nigeria": "NG", "tchad": "TD", "centrafrique": "CF",
    "republique_centrafricaine": "CF", "cameroun": "CM", "guinee_equatoriale": "GQ", "gabon": "GA", "congo": "CG",
    "republique_du_congo": "CG", "congo_brazzaville": "CG", "rd_congo": "CD", "rdc": "CD",
    "republique_democratique_du_congo": "CD", "guinee_bissau": "GW", "maroc": "MA", "tunisie": "TN", "france": "FR",
}


def codes_zones(zones):
    """Code à deux lettres de chaque zone : code pays connu, sinon les premières lettres (sans doublon)."""
    codes, pris = {}, set()
    for z in zones:
        s = slug(z, "").upper()
        c = CODES_PAYS.get(slug(z)) or s[:2]
        i = 2
        while c in pris and i < len(s):
            c, i = s[0] + s[i], i + 1
        codes[z] = c
        pris.add(c)
    return codes


def identifiant(code_zone, rang):
    return f"ENQ{code_zone}{rang:02d}"


CAPITALES = {
    "senegal": "Dakar", "mali": "Bamako", "cote_d_ivoire": "Abidjan", "burkina_faso": "Ouagadougou", "niger": "Niamey",
    "togo": "Lomé", "benin": "Cotonou", "guinee": "Conakry", "guinee_bissau": "Bissau", "cameroun": "Douala",
    "gabon": "Libreville", "congo": "Brazzaville", "rd_congo": "Kinshasa", "tchad": "N'Djamena", "mauritanie": "Nouakchott",
}
MAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.I)


def exemples(zones):
    """Deux lignes fictives montrant le remplissage attendu (pays de l'étude)."""
    z1, z2 = (zones + zones)[:2] if zones else ("Sénégal", "Côte d'Ivoire")
    # numéro local fictif à la longueur du pays (9 chiffres au Sénégal, 10 en Côte d'Ivoire, 8 le plus souvent)
    local = {"senegal": "771234567", "cote_d_ivoire": "0701020304", "cameroun": "677123456", "rd_congo": "812345678", "guinee": "621123456"}
    tel = lambda z, defaut: f"+{INDICATIFS.get(slug(z), '221')} {local.get(slug(z), defaut)}"
    return [
        {"nom": "DIOP", "prenom": "Awa", "pays": z1, "ville": CAPITALES.get(slug(z1), z1), "telephone": tel(z1, "90123456"), "mail": "awa.diop@exemple.com"},
        {"nom": "KOUASSI", "prenom": "Jean-Marc", "pays": z2, "ville": CAPITALES.get(slug(z2), z2), "telephone": tel(z2, "70123456"), "mail": "jm.kouassi@exemple.com"},
    ]


def modele_xlsx(chemin, zones, besoin):
    """Fichier à renseigner par les RH : feuille Enquêteurs (vide, pays en liste déroulante), consignes et exemple."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Enquêteurs"
    _feuille(ws, [c for c, _, _ in COLONNES], [], {c: 24 for c, _, _ in COLONNES})
    ws.column_dimensions["F"].width = 34
    nb = max(sum(besoin.values()) + 20, 50)
    for row in ws.iter_rows(min_row=2, max_row=nb + 1, min_col=5, max_col=5):
        row[0].number_format = "@"  # le « + » de l'indicatif est conservé
    if zones and len(",".join(zones)) < 250:
        dv = DataValidation(type="list", formula1='"' + ",".join(z.replace('"', "") for z in zones) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Choisissez un pays de l'étude.", "Pays de résidence"
        ws.add_data_validation(dv)
        dv.add(f"C2:C{nb + 1}")
    ws2 = wb.create_sheet("Consignes")
    _feuille(ws2, ["Colonne", "Consigne"], [[c, t] for c, _, t in COLONNES]
             + [["", ""], ["Une ligne", "Une ligne par enquêteur retenu, sans ligne vide ni cellule fusionnée ; ne pas renommer les colonnes."],
                ["Identifiant", "Ne pas en ajouter : Datatym attribue l'identifiant de chaque enquêteur à l'import "
                 "(ENQ, code du pays et numéro d'ordre : ENQML01, ENQCI01, ENQSN01…)."]],
             {"Colonne": 20, "Consigne": 90})
    ws2.append([])
    ws2.append(["Exemple de remplissage (personnes fictives)"])
    ws2.cell(ws2.max_row, 1).font = Font(bold=True)
    for e in exemples(zones):
        ws2.append([e[k] for _, k, _ in COLONNES])
    ws3 = wb.create_sheet("Besoin")
    _feuille(ws3, ["Pays", "Enquêteurs à recruter"], [[z, besoin.get(z, 0)] for z in zones] + [["Total", sum(besoin.values())]], {"Pays": 28})
    for c in ws3[ws3.max_row]:
        c.font = Font(bold=True)
        c.alignment = Alignment()
    wb.properties.creator = wb.properties.lastModifiedBy = "DATATYM"
    wb.save(chemin)
    return chemin


def _cle_colonne(entete):
    s = slug(entete or "")
    if s.startswith("prenom"):
        return "prenom"
    if s == "nom" or s.startswith("nom_de_famille") or s.startswith("nom_"):
        return "nom"
    if s.startswith("pays"):
        return "pays"
    if s.startswith("ville") or s.startswith("localite"):
        return "ville"
    if s.startswith(("tel", "telephone", "numero", "contact", "whatsapp", "portable")):
        return "telephone"
    if "mail" in s or s.startswith("courriel"):
        return "mail"
    return None


def _lignes(chemin):
    if str(chemin).lower().endswith(".csv"):
        texte = open(chemin, encoding="utf-8-sig", errors="replace").read()
        sep = ";" if texte.count(";") > texte.count(",") else ","
        return list(csv.reader(texte.splitlines(), delimiter=sep))
    ws = load_workbook(chemin, read_only=True, data_only=True).worksheets[0]
    return [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]


def zone_de(pays, ville, zones):
    """Zone de l'étude correspondant au pays (ou à la ville, quand les zones sont des villes)."""
    par = {slug(z): z for z in zones}
    for v in (pays, ville):
        if slug(v) in par:
            return par[slug(v)]
    return zones[0] if len(zones) == 1 else None


def telephone(brut, zone):
    """Numéro au format +indicatif ; l'indicatif du pays est ajouté s'il manque. Renvoie (numéro, remarque)."""
    t = str(brut or "").strip()
    if t.endswith(".0"):  # nombre lu depuis Excel
        t = t[:-2]
    chiffres = re.sub(r"\D", "", t)
    if not chiffres:
        return "", "téléphone manquant"
    if t.startswith("+"):
        pass
    elif chiffres.startswith("00"):
        chiffres = chiffres[2:]
    else:
        ind = INDICATIFS.get(slug(zone or ""))
        if ind and chiffres.startswith(ind) and len(chiffres) > len(ind) + 7:
            pass
        elif ind:
            return f"+{ind} {chiffres}", "indicatif ajouté"
        else:
            return f"+{chiffres}", "indicatif à vérifier"
    if not 8 <= len(chiffres) <= 15:
        return "+" + chiffres, "numéro de longueur inhabituelle"
    ind = next((i for i in sorted(set(INDICATIFS.values()), key=len, reverse=True) if chiffres.startswith(i)), None)
    return (f"+{ind} {chiffres[len(ind):]}" if ind else f"+{chiffres}"), None


def lire(chemin, zones, besoin):
    """Liste des RH -> (équipe, alertes, rejetés). Une ligne sans nom ou hors des pays de l'étude est rejetée."""
    lignes = [l for l in _lignes(chemin) if any(str(x).strip() for x in l)]
    if not lignes:
        raise ValueError("Le fichier est vide.")
    i_ent = next((i for i, l in enumerate(lignes[:10]) if sum(1 for x in l if _cle_colonne(x)) >= 3), None)
    if i_ent is None:
        raise ValueError("Colonnes introuvables : le fichier doit contenir Nom, Prénom, Pays de résidence, Ville, Téléphone et Mail.")
    cols = {}
    for j, x in enumerate(lignes[i_ent]):
        k = _cle_colonne(x)
        if k and k not in cols:
            cols[k] = j
    manquantes = [c for c, k, _ in COLONNES if k not in cols]
    exemple = {(e["nom"], e["prenom"]) for e in exemples(zones)}
    equipe, alertes, rejetes, vus, completes = [], [], [], {}, []
    for n, l in enumerate(lignes[i_ent + 1:], i_ent + 2):
        v = {k: str(l[j]).strip() if j < len(l) else "" for k, j in cols.items()}
        if not (v.get("nom") or v.get("prenom")):
            continue
        if (v.get("nom", "").upper(), v.get("prenom", "")) in exemple:
            continue  # lignes d'exemple recopiées
        qui = f"Ligne {n} ({v.get('prenom', '')} {v.get('nom', '')})".replace("( ", "(").replace(" )", ")")
        zone = zone_de(v.get("pays", ""), v.get("ville", ""), zones)
        if not zone:
            rejetes.append(f"{qui} : pays « {v.get('pays', '')} » hors de l'étude")
            continue
        tel, rem = telephone(v.get("telephone"), zone)
        if rem == "indicatif ajouté":
            completes.append(str(n))
        elif rem:
            alertes.append(f"{qui} : {rem}")
        mail = v.get("mail", "").lower()
        if not mail:
            alertes.append(f"{qui} : mail manquant")
        elif not MAIL.match(mail):
            alertes.append(f"{qui} : mail invalide « {mail} »")
        for cle, val in (("téléphone", tel), ("mail", mail)):
            if val and (cle, val) in vus:
                alertes.append(f"{qui} : même {cle} que la ligne {vus[(cle, val)]}")
            elif val:
                vus[(cle, val)] = n
        equipe.append({"nom": v.get("nom", "").upper(), "prenom": v.get("prenom", ""), "pays": zone,
                       "ville": v.get("ville", ""), "telephone": tel, "mail": mail})
    if not equipe:
        raise ValueError("Aucun enquêteur lu dans le fichier." + (f" Lignes rejetées : {'; '.join(rejetes[:3])}." if rejetes else ""))
    if manquantes:
        alertes.insert(0, "Colonnes absentes : " + ", ".join(manquantes))
    if completes:
        alertes.append(f"Indicatif du pays ajouté à {len(completes)} numéro(s) (ligne{'s' if len(completes) > 1 else ''} {', '.join(completes)})")
    for z in zones:
        nb = sum(1 for e in equipe if e["pays"] == z)
        if besoin.get(z) and nb != besoin[z]:
            alertes.append(f"{z} : {nb} enquêteur(s) pour {besoin[z]} prévu(s)")
    return equipe, alertes, rejetes
